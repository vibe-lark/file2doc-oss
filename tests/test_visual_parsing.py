from __future__ import annotations

import io
import json
from pathlib import Path
from types import SimpleNamespace
import zipfile

from docx import Document
from fastapi.testclient import TestClient
from markitdown import StreamInfo
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from file2doc.app import create_app
from file2doc.parsers import (
    ParsedContent,
    ParseOptions,
    _warnings_from_markdown,
    parse_content_markdown,
)
from file2doc_markitdown_visual.plugin import (
    VISUAL_RESULT_SCHEMA,
    VisualArtifactCollector,
    VisualDiagnosticArtifact,
    VisualImageConverter,
)


class FakeResponses:
    def __init__(self, results, *, outputs=None):
        self.results = list(results)
        self.outputs = list(outputs or [[] for _ in self.results])
        self.calls: list[dict] = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        result = self.results.pop(0)
        if isinstance(result, Exception):
            raise result
        return SimpleNamespace(
            output_text=json.dumps(result),
            output=self.outputs.pop(0),
        )


def visual_payload(
    description: str = "A red status icon.",
    visible_text: list[str] | None = None,
) -> dict:
    return {
        "description": description,
        "visible_text": visible_text or ["OK"],
        "layout": "The icon is centered.",
        "warnings": [],
    }


def png_bytes(color: str = "red") -> bytes:
    output = io.BytesIO()
    Image.new("RGB", (24, 24), color).save(output, format="PNG")
    return output.getvalue()


def docx_bytes(images: list[bytes], *, text: str = "Native text") -> bytes:
    document = Document()
    document.add_paragraph(text)
    for image in images:
        document.add_picture(io.BytesIO(image))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pptx_bytes(image: bytes) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(io.BytesIO(image), Inches(1), Inches(1))
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def xlsx_bytes(image: bytes) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Evidence"
    sheet["A1"] = "Native cell"
    sheet.add_image(SpreadsheetImage(io.BytesIO(image)), "B2")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def pdf_bytes(*, image: bytes | None = None, native_text: str | None = None) -> bytes:
    output = io.BytesIO()
    document = canvas.Canvas(output, pagesize=(300, 300))
    if native_text:
        document.drawString(20, 270, native_text)
    if image is not None:
        document.drawImage(ImageReader(io.BytesIO(image)), 20, 120, 60, 60)
    document.showPage()
    document.save()
    return output.getvalue()


def native_text_pdf_bytes(*pages: str) -> bytes:
    output = io.BytesIO()
    document = canvas.Canvas(output, pagesize=(300, 300))
    for text in pages:
        document.drawString(20, 270, text)
        document.showPage()
    document.save()
    return output.getvalue()


def terminal_job(client: TestClient, poll_url: str) -> dict:
    for _ in range(100):
        job = client.get(poll_url).json()
        if job["status"] in {"completed", "completed_with_warnings", "failed"}:
            return job
    raise AssertionError("parse job did not reach a terminal state")


def test_visual_provider_uses_generic_schema_and_image_process_tools() -> None:
    responses = FakeResponses(
        [visual_payload()],
        outputs=[
            [
                {
                    "type": "image_process",
                    "status": "completed",
                    "action": {
                        "type": "zoom",
                        "arguments": {"bbox_str": "0,0,12,12"},
                        "result_image_url": "https://provider.invalid/cropped.png",
                    },
                }
            ]
        ],
    )
    collector = VisualArtifactCollector()
    converter = VisualImageConverter(
        client=SimpleNamespace(responses=responses),
        model="ep-visual",
        artifact_collector=collector,
    )

    result = converter.convert(
        io.BytesIO(png_bytes()),
        StreamInfo(filename="status.png", mimetype="image/png"),
        visual_location="DOCX image 1",
    )

    request = responses.calls[0]
    assert request["text"]["format"]["schema"] == VISUAL_RESULT_SCHEMA
    assert request["tools"] == [
        {
            "type": "image_process",
            "point": {"type": "disabled"},
            "grounding": {"type": "disabled"},
            "zoom": {"type": "enabled"},
            "rotate": {"type": "enabled"},
        }
    ]
    assert "candidateNumericValues" not in json.dumps(request)
    assert "business fields" in request["input"][0]["content"][1]["text"]
    assert "Visual Description" in result.markdown
    assert "Visual Parsing Runtime" not in result.markdown
    assert collector.results[0].locations == ("DOCX image 1",)
    assert collector.results[0].image_process_audits[0].action_type == "zoom"
    assert collector.results[0].image_process_audits[0].diagnostic_ref is None
    assert collector.artifacts == ()


def test_docx_parses_every_image_and_deduplicates_exact_bytes(tmp_path: Path) -> None:
    duplicate = png_bytes()
    source = tmp_path / "images.docx"
    source.write_bytes(docx_bytes([duplicate, duplicate]))
    responses = FakeResponses([visual_payload()])

    parsed = parse_content_markdown(
        source,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ParseOptions(
            visual_client=SimpleNamespace(responses=responses),
            visual_model="ep-visual",
        ),
    )

    assert len(responses.calls) == 1
    assert len(parsed.visual_sources) == 1
    assert parsed.visual_sources[0].locations == ("DOCX image 1", "DOCX image 2")
    assert parsed.visual_results[0].locations == ("DOCX image 1", "DOCX image 2")
    assert parsed.markdown.count("A red status icon.") == 2


def test_pptx_and_xlsx_parse_embedded_images_without_page_context(
    tmp_path: Path,
) -> None:
    source_image = png_bytes()
    cases = [
        (
            "slides.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            pptx_bytes(source_image),
            "PPTX slide 1",
        ),
        (
            "sheet.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            xlsx_bytes(source_image),
            "XLSX sheet Evidence, cell B2",
        ),
    ]
    for filename, content_type, source_bytes, expected_location in cases:
        source = tmp_path / filename
        source.write_bytes(source_bytes)
        responses = FakeResponses([visual_payload(description=f"Parsed {filename} image.")])

        parsed = parse_content_markdown(
            source,
            content_type,
            ParseOptions(
                visual_client=SimpleNamespace(responses=responses),
                visual_model="ep-visual",
            ),
        )

        assert len(responses.calls) == 1
        assert expected_location in parsed.visual_results[0].locations[0]
        assert f"Parsed {filename} image." in parsed.markdown


def test_pdf_uses_native_text_and_vlm_regions_but_publishes_complete_pages(
    tmp_path: Path,
) -> None:
    embedded = tmp_path / "embedded.pdf"
    embedded.write_bytes(pdf_bytes(image=png_bytes(), native_text="Native PDF text"))
    embedded_responses = FakeResponses([visual_payload(description="Complete page.")])

    embedded_result = parse_content_markdown(
        embedded,
        "application/pdf",
        ParseOptions(
            visual_client=SimpleNamespace(responses=embedded_responses),
            visual_model="ep-visual",
        ),
    )

    assert "Native PDF text" in embedded_result.markdown
    assert "Complete page." in embedded_result.markdown
    assert embedded_result.visual_results[0].locations == ("page 1 image 1",)
    assert len(embedded_result.visual_sources) == 1
    with Image.open(io.BytesIO(embedded_result.visual_sources[0].content)) as image:
        assert image.size == (24, 24)
    source_hash = embedded_result.visual_sources[0].content_sha256
    assert "![page-1-image](images/pages/page_001.png)" in embedded_result.markdown
    assert f"images/visual/{source_hash}.png" not in embedded_result.markdown

    scanned = tmp_path / "scanned.pdf"
    scanned.write_bytes(pdf_bytes())
    scanned_responses = FakeResponses([visual_payload(description="Scanned page.")])
    scanned_result = parse_content_markdown(
        scanned,
        "application/pdf",
        ParseOptions(
            visual_client=SimpleNamespace(responses=scanned_responses),
            visual_model="ep-visual",
        ),
    )

    assert "Scanned page." in scanned_result.markdown
    assert scanned_result.visual_results[0].locations == ("page 1",)


def test_pdf_package_excludes_internal_visual_region_images(tmp_path: Path) -> None:
    responses = FakeResponses([visual_payload(description="Embedded diagram.")])
    client = TestClient(
        create_app(
            storage_root=tmp_path / "storage",
            auth_enabled=False,
            parse_options=ParseOptions(
                visual_client=SimpleNamespace(responses=responses),
                visual_model="ep-visual",
            ),
        )
    )
    created = client.post(
        "/parse-jobs/upload",
        files={
            "file": (
                "mixed.pdf",
                pdf_bytes(image=png_bytes(), native_text="Native PDF text"),
                "application/pdf",
            )
        },
    ).json()

    job = terminal_job(client, created["poll_url"])
    manifest = client.get(job["result"]["manifest_url"]).json()
    content = client.get(job["result"]["content_url"]).text

    assert [item["kind"] for item in manifest["media_index"]] == [
        "page_image",
        "thumbnail",
    ]
    assert not any(
        item["kind"] == "visual_source_image" for item in manifest["artifacts"]
    )
    assert "![page-1-image](images/pages/page_001.png)" in content
    assert "### Visual region: page 1 image 1" in content


def test_pdf_with_extractable_text_does_not_call_visual_provider(tmp_path: Path) -> None:
    source = tmp_path / "native.pdf"
    source.write_bytes(pdf_bytes(native_text="Native PDF text"))
    responses = FakeResponses([])

    parsed = parse_content_markdown(
        source,
        "application/pdf",
        ParseOptions(
            visual_client=SimpleNamespace(responses=responses),
            visual_model="ep-visual",
        ),
    )

    assert responses.calls == []
    assert "Native PDF text" in parsed.markdown
    assert "![page-1-image](images/pages/page_001.png)" in parsed.markdown


def test_multi_page_pdf_publishes_every_complete_page(tmp_path: Path) -> None:
    responses = FakeResponses([])
    client = TestClient(
        create_app(
            storage_root=tmp_path / "storage",
            auth_enabled=False,
            parse_options=ParseOptions(
                visual_client=SimpleNamespace(responses=responses),
                visual_model="ep-visual",
            ),
        )
    )
    created = client.post(
        "/parse-jobs/upload",
        files={
            "file": (
                "native-pages.pdf",
                native_text_pdf_bytes("Page one", "Page two"),
                "application/pdf",
            )
        },
    ).json()

    job = terminal_job(client, created["poll_url"])
    manifest = client.get(job["result"]["manifest_url"]).json()
    content = client.get(job["result"]["content_url"]).text

    assert responses.calls == []
    assert len(manifest["page_index"]) == 2
    assert [
        item["path"]
        for item in manifest["media_index"]
        if item["kind"] == "page_image"
    ] == ["images/pages/page_001.png", "images/pages/page_002.png"]
    assert "![page-1-image](images/pages/page_001.png)" in content
    assert "![page-2-image](images/pages/page_002.png)" in content


def test_scanned_pdf_visual_failure_returns_warning_and_empty_result(
    tmp_path: Path,
) -> None:
    source = tmp_path / "scanned.pdf"
    source.write_bytes(pdf_bytes())

    parsed = parse_content_markdown(
        source,
        "application/pdf",
        ParseOptions(
            visual_client=SimpleNamespace(
                responses=FakeResponses([RuntimeError("provider unavailable")])
            ),
            visual_model="ep-visual",
        ),
    )

    assert parsed.markdown == ""
    assert parsed.diagnostics["empty_result"] is True
    assert parsed.warnings[0]["code"] == "visual_item_failed"
    assert parsed.visual_sources[0].locations == ("page 1",)


def test_visual_warning_extraction_stops_before_following_native_text() -> None:
    warnings = _warnings_from_markdown(
        """### Visual Warnings

- Small text is blurred.
- One label is unreadable.

Native PDF text after the embedded image must not become a warning.
"""
    )

    assert warnings == [
        {
            "severity": "warning",
            "code": "visual_item_warning",
            "message": "Small text is blurred. One label is unreadable.",
        }
    ]


def test_result_package_contains_original_image_json_and_inline_markdown(
    tmp_path: Path,
) -> None:
    responses = FakeResponses([visual_payload()])
    client = TestClient(
        create_app(
            storage_root=tmp_path / "storage",
            auth_enabled=False,
            parse_options=ParseOptions(
                visual_client=SimpleNamespace(responses=responses),
                visual_model="ep-visual",
            ),
        )
    )
    created = client.post(
        "/parse-jobs/upload",
        files={
            "file": (
                "visual.docx",
                docx_bytes([png_bytes()]),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    ).json()

    job = terminal_job(client, created["poll_url"])
    assert job["status"] == "completed"
    manifest = client.get(job["result"]["manifest_url"]).json()
    media = next(item for item in manifest["media_index"] if item["kind"] == "embedded_image")
    assert media["visual_parse_status"] == "completed"
    assert media["visual_result_artifact_id"]
    artifacts = {item["artifact_id"]: item for item in manifest["artifacts"]}
    assert artifacts[media["artifact_id"]]["kind"] == "visual_source_image"
    visual_artifact = artifacts[media["visual_result_artifact_id"]]
    assert visual_artifact["kind"] == "visual_parse_result"
    visual_result = client.get(
        f"/parse-jobs/{job['job_id']}/artifacts/{visual_artifact['artifact_id']}"
    ).json()
    assert visual_result["description"] == "A red status icon."
    assert visual_result["visible_text"] == ["OK"]
    assert visual_result["provider"] == "ark-responses"
    content = client.get(job["result"]["content_url"]).text
    assert f"![{media['id']}]({media['path']})" in content
    assert "A red status icon." in content


def test_visual_failure_preserves_native_content_and_original_image(
    tmp_path: Path,
) -> None:
    responses = FakeResponses([RuntimeError("provider unavailable")])
    client = TestClient(
        create_app(
            storage_root=tmp_path / "storage",
            auth_enabled=False,
            parse_options=ParseOptions(
                visual_client=SimpleNamespace(responses=responses),
                visual_model="ep-visual",
            ),
        )
    )
    created = client.post(
        "/parse-jobs/upload",
        files={
            "file": (
                "visual.docx",
                docx_bytes([png_bytes()], text="Keep this paragraph"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    ).json()

    job = terminal_job(client, created["poll_url"])
    assert job["status"] == "completed_with_warnings"
    assert job["error"] is None
    manifest = client.get(job["result"]["manifest_url"]).json()
    assert manifest["warnings"][0]["code"] == "visual_item_warning"
    media = next(item for item in manifest["media_index"] if item["kind"] == "embedded_image")
    assert media["visual_parse_status"] == "warning"
    assert media["visual_result_artifact_id"] is None
    assert client.get(job["result"]["content_url"]).text.startswith("Keep this paragraph")


def test_image_process_diagnostic_artifacts_are_not_in_result_package(
    tmp_path: Path,
    monkeypatch,
) -> None:
    import file2doc.store as store_module

    diagnostic = VisualDiagnosticArtifact(
        kind="visual_image_process_result",
        source_ref="visual-source-1",
        diagnostic_ref="visual-source-1:image_process:0001",
        media_type="image/png",
        content=png_bytes(),
        action_type="zoom",
        arguments=(("region", "center"),),
        status="completed",
        warnings=(),
    )
    monkeypatch.setattr(
        store_module,
        "parse_content_markdown",
        lambda source_path, content_type, options=None: ParsedContent(
            markdown="Native text",
            diagnostics={"parser": "test"},
            visual_artifacts=(diagnostic,),
        ),
    )

    with TestClient(create_app(storage_root=tmp_path, auth_enabled=False)) as client:
        created = client.post(
            "/parse-jobs/upload",
            files={"file": ("visual.txt", b"Native text", "text/plain")},
        ).json()
        job = terminal_job(client, created["poll_url"])
        manifest = client.get(job["result"]["manifest_url"]).json()
        assert not any(item.get("diagnostic_only") for item in manifest["artifacts"])
        package = client.get(job["result"]["package_url"])
        assert package.status_code == 200
        with zipfile.ZipFile(io.BytesIO(package.content)) as archive:
            assert not any(
                name.startswith("diagnostics/image-process/")
                for name in archive.namelist()
            )


def test_visual_configuration_uses_only_visual_environment(monkeypatch) -> None:
    monkeypatch.setenv("FILE2DOC_VISUAL_MODEL", "ep-visual")
    monkeypatch.setenv("FILE2DOC_VISUAL_API_KEY", "secret")
    monkeypatch.setenv("FILE2DOC_VISUAL_BASE_URL", "https://ark.example.test/api/v3")
    monkeypatch.setenv("FILE2DOC_VISUAL_ITEM_TIMEOUT_SECONDS", "12")
    monkeypatch.setenv("FILE2DOC_VISUAL_JOB_DEADLINE_SECONDS", "34")
    monkeypatch.setenv("FILE2DOC_VISUAL_MAX_CONCURRENCY", "2")
    monkeypatch.setenv("FILE2DOC_OCR_MODEL", "legacy-must-not-be-used")

    options = ParseOptions.from_env()

    assert options.visual_configured is True
    assert options.visual_model == "ep-visual"
    assert options.visual_item_timeout_seconds == 12
    assert options.visual_job_deadline_seconds == 34
    assert options.visual_max_concurrency == 2
