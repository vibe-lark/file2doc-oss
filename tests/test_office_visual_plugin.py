import io
import json
from types import SimpleNamespace

import pytest
from docx import Document
from markitdown import MarkItDown, StreamInfo
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from file2doc_markitdown_visual.embedded import EmbeddedVisualParser


class _VisualClient:
    def __init__(self, results):
        self._results = iter(results)
        self.responses = self
        self.calls = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        result = next(self._results)
        if isinstance(result, Exception):
            raise result
        return SimpleNamespace(output_text=json.dumps(result))


def test_docx_embedded_image_is_parsed_at_its_document_position():
    document = Document()
    document.add_paragraph("Before embedded image")
    document.add_picture(io.BytesIO(_png_bytes()))
    document.add_paragraph("After embedded image")
    source = io.BytesIO()
    document.save(source)
    source.seek(0)

    client = _VisualClient(
        [_visual_result(description="A blue laboratory logo [alpha_beta].")]
    )
    markitdown = MarkItDown(
        enable_plugins=True,
        visual_client=client,
        visual_model="ep-visual",
    )

    markdown = markitdown.convert_stream(
        source,
        stream_info=StreamInfo(extension=".docx"),
    ).markdown

    assert markdown.index("Before embedded image") < markdown.index(
        "A blue laboratory logo [alpha_beta]."
    )
    assert markdown.index("A blue laboratory logo [alpha_beta].") < markdown.index(
        "After embedded image"
    )
    assert "## Visible Text\n\nNone." in markdown
    assert len(client.calls) == 1


def test_docx_reuses_visual_result_for_exact_duplicate_pixel_bytes():
    duplicate = _png_bytes()
    document = Document()
    document.add_paragraph("First")
    document.add_picture(io.BytesIO(duplicate))
    document.add_paragraph("Between")
    document.add_picture(io.BytesIO(duplicate))
    document.add_paragraph("Last")
    source = io.BytesIO()
    document.save(source)
    source.seek(0)

    client = _VisualClient([_visual_result(description="Repeated logo.")])
    markdown = MarkItDown(
        enable_plugins=True,
        visual_client=client,
        visual_model="ep-visual",
    ).convert_stream(source, stream_info=StreamInfo(extension=".docx")).markdown

    assert markdown.count("Repeated logo.") == 2
    assert markdown.index("First") < markdown.index("DOCX image 1")
    assert markdown.index("DOCX image 1") < markdown.index("Between")
    assert markdown.index("Between") < markdown.index("DOCX image 2")
    assert markdown.index("DOCX image 2") < markdown.index("Last")
    assert len(client.calls) == 1


@pytest.mark.parametrize(
    "extension",
    [".docx", ".pptx", ".xlsx"],
)
def test_same_markitdown_instance_starts_fresh_cache_for_each_office_document(
    extension,
):
    duplicate = _png_bytes()
    document_builder = {
        ".docx": _docx_with_images,
        ".pptx": _pptx_with_images,
        ".xlsx": _xlsx_with_images,
    }[extension]
    client = _VisualClient(
        [
            _visual_result(description="First document result."),
            _visual_result(description="Second document result."),
        ]
    )
    markitdown = MarkItDown(
        enable_plugins=True,
        visual_client=client,
        visual_model="ep-visual",
    )

    first = markitdown.convert_stream(
        document_builder([duplicate, duplicate]),
        stream_info=StreamInfo(extension=extension),
    ).markdown
    second = markitdown.convert_stream(
        document_builder([duplicate, duplicate]),
        stream_info=StreamInfo(extension=extension),
    ).markdown

    assert first.count("First document result.") == 2
    assert second.count("Second document result.") == 2
    assert len(client.calls) == 2


def test_docx_native_text_cannot_collide_with_generated_image_placeholder():
    document = Document()
    document.add_paragraph("FILE2DOCVISUALBLOCK0")
    document.add_picture(io.BytesIO(_png_bytes()))
    source = io.BytesIO()
    document.save(source)
    source.seek(0)

    client = _VisualClient([_visual_result(description="Only the real image.")])
    markdown = MarkItDown(
        enable_plugins=True,
        visual_client=client,
        visual_model="ep-visual",
    ).convert_stream(source, stream_info=StreamInfo(extension=".docx")).markdown

    assert "FILE2DOCVISUALBLOCK0" in markdown
    assert markdown.count("Only the real image.") == 1


def test_docx_failed_image_is_an_explicit_warning_and_other_content_survives():
    document = Document()
    document.add_paragraph("Native content before")
    document.add_picture(io.BytesIO(_png_bytes(color=(200, 10, 10))))
    document.add_paragraph("Native content between")
    document.add_picture(io.BytesIO(_png_bytes(color=(10, 200, 10))))
    document.add_paragraph("Native content after")
    source = io.BytesIO()
    document.save(source)
    source.seek(0)

    client = _VisualClient(
        [
            RuntimeError("provider unavailable"),
            _visual_result(description="Successful second image."),
        ]
    )
    markdown = MarkItDown(
        enable_plugins=True,
        visual_client=client,
        visual_model="ep-visual",
    ).convert_stream(source, stream_info=StreamInfo(extension=".docx")).markdown

    assert (
        "> Warning: visual parsing failed at DOCX image 1: "
        "RuntimeError: provider unavailable"
    ) in markdown
    assert "Successful second image." in markdown
    assert "Native content before" in markdown
    assert "Native content between" in markdown
    assert "Native content after" in markdown
    assert len(client.calls) == 2


def test_pptx_picture_has_description_and_visible_text_in_shape_order():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    before = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(0.5))
    before.text = "Before picture"
    slide.shapes.add_picture(io.BytesIO(_png_bytes()), Inches(1), Inches(2))
    after = slide.shapes.add_textbox(Inches(1), Inches(4), Inches(4), Inches(0.5))
    after.text = "After picture"
    source = io.BytesIO()
    presentation.save(source)
    source.seek(0)

    client = _VisualClient(
        [
            _visual_result(
                description="A gauge display.",
                visible_text=["H 88.52"],
            )
        ]
    )
    markdown = MarkItDown(
        enable_plugins=True,
        visual_client=client,
        visual_model="ep-visual",
    ).convert_stream(source, stream_info=StreamInfo(extension=".pptx")).markdown

    assert "<!-- Slide number: 1 -->" in markdown
    assert markdown.index("Before picture") < markdown.index("A gauge display.")
    assert markdown.index("A gauge display.") < markdown.index("H 88.52")
    assert markdown.index("H 88.52") < markdown.index("After picture")
    assert len(client.calls) == 1


def test_pptx_failed_picture_warning_keeps_other_shapes_and_visuals():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(0.5), Inches(4), Inches(0.5)).text = (
        "Native heading"
    )
    slide.shapes.add_picture(
        io.BytesIO(_png_bytes(color=(150, 0, 0))), Inches(1), Inches(1.5)
    )
    slide.shapes.add_picture(
        io.BytesIO(_png_bytes(color=(0, 150, 0))), Inches(1), Inches(3)
    )
    slide.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(0.5)).text = (
        "Native ending"
    )
    source = io.BytesIO()
    presentation.save(source)
    source.seek(0)

    client = _VisualClient(
        [
            RuntimeError("visual timeout"),
            _visual_result(description="Second picture survives."),
        ]
    )
    markdown = MarkItDown(
        enable_plugins=True,
        visual_client=client,
        visual_model="ep-visual",
    ).convert_stream(source, stream_info=StreamInfo(extension=".pptx")).markdown

    assert "Native heading" in markdown
    assert "visual parsing failed at PPTX slide 1" in markdown
    assert "RuntimeError: visual timeout" in markdown
    assert "Second picture survives." in markdown
    assert "Native ending" in markdown
    assert len(client.calls) == 2


def test_xlsx_picture_keeps_sheet_and_anchor_cell_context():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet.append(["Metric", "Value"])
    sheet.append(["Haze", "88.52"])
    sheet.add_image(SpreadsheetImage(io.BytesIO(_png_bytes())), "B4")
    source = io.BytesIO()
    workbook.save(source)
    source.seek(0)

    client = _VisualClient(
        [_visual_result(description="A small instrument screenshot.")]
    )
    markdown = MarkItDown(
        enable_plugins=True,
        visual_client=client,
        visual_model="ep-visual",
    ).convert_stream(source, stream_info=StreamInfo(extension=".xlsx")).markdown

    assert "## Summary" in markdown
    assert "| Haze | 88.52 |" in markdown
    assert "XLSX sheet Summary, cell B4" in markdown
    assert "A small instrument screenshot." in markdown
    assert markdown.index("| Haze | 88.52 |") < markdown.index(
        "XLSX sheet Summary, cell B4"
    )
    assert len(client.calls) == 1


def test_xlsx_duplicate_images_across_sheets_share_one_visual_request():
    duplicate = _png_bytes(color=(90, 90, 90))
    workbook = Workbook()
    first = workbook.active
    first.title = "First"
    first.append(["Native", "one"])
    first.add_image(SpreadsheetImage(io.BytesIO(duplicate)), "A3")
    second = workbook.create_sheet("Second")
    second.append(["Native", "two"])
    second.add_image(SpreadsheetImage(io.BytesIO(duplicate)), "C5")
    source = io.BytesIO()
    workbook.save(source)
    source.seek(0)

    client = _VisualClient([_visual_result(description="Shared company logo.")])
    markdown = MarkItDown(
        enable_plugins=True,
        visual_client=client,
        visual_model="ep-visual",
    ).convert_stream(source, stream_info=StreamInfo(extension=".xlsx")).markdown

    assert markdown.index("## First") < markdown.index("XLSX sheet First, cell A3")
    assert markdown.index("XLSX sheet First, cell A3") < markdown.index("## Second")
    assert markdown.index("## Second") < markdown.index("XLSX sheet Second, cell C5")
    assert markdown.count("Shared company logo.") == 2
    assert len(client.calls) == 1


def test_gif_embedded_asset_is_normalized_to_png_before_visual_request():
    document = Document()
    document.add_picture(io.BytesIO(_gif_bytes()))
    source = io.BytesIO()
    document.save(source)
    source.seek(0)

    client = _VisualClient([_visual_result(description="Animated status icon.")])
    markdown = MarkItDown(
        enable_plugins=True,
        visual_client=client,
        visual_model="ep-visual",
    ).convert_stream(source, stream_info=StreamInfo(extension=".docx")).markdown

    image_url = client.calls[0]["input"][0]["content"][0]["image_url"]
    assert image_url.startswith("data:image/png;base64,")
    assert "Animated status icon." in markdown


def test_vector_embedded_asset_returns_warning_without_calling_provider():
    client = _VisualClient([])
    parser = EmbeddedVisualParser(client=client, model="ep-visual")

    result = parser.new_session().parse(
        b'<svg xmlns="http://www.w3.org/2000/svg"><circle r="4"/></svg>',
        content_type="image/svg+xml",
    )

    assert result.warning is not None
    assert "unsupported non-raster" in result.warning.lower()
    assert client.calls == []


def _visual_result(*, description, visible_text=None):
    return {
        "description": description,
        "visibleText": visible_text or [],
        "candidateNumericValues": [],
        "layout": "Centered image.",
        "warnings": [],
    }


def _png_bytes(color=(32, 64, 192)):
    image = Image.new("RGB", (24, 16), color=color)
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    return buffer.getvalue()


def _gif_bytes():
    first = Image.new("RGB", (24, 16), color=(255, 0, 0))
    second = Image.new("RGB", (24, 16), color=(0, 255, 0))
    buffer = io.BytesIO()
    first.save(buffer, format="GIF", save_all=True, append_images=[second])
    return buffer.getvalue()


def _docx_with_images(images):
    document = Document()
    for image_bytes in images:
        document.add_picture(io.BytesIO(image_bytes))
    source = io.BytesIO()
    document.save(source)
    source.seek(0)
    return source


def _pptx_with_images(images):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for index, image_bytes in enumerate(images):
        slide.shapes.add_picture(
            io.BytesIO(image_bytes),
            Inches(1),
            Inches(1 + index * 2),
        )
    source = io.BytesIO()
    presentation.save(source)
    source.seek(0)
    return source


def _xlsx_with_images(images):
    workbook = Workbook()
    sheet = workbook.active
    for index, image_bytes in enumerate(images, start=1):
        sheet.add_image(
            SpreadsheetImage(io.BytesIO(image_bytes)),
            f"A{index}",
        )
    source = io.BytesIO()
    workbook.save(source)
    source.seek(0)
    return source
