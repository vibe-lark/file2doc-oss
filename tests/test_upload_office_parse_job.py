import io
from pathlib import Path
import json
from types import SimpleNamespace

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SpreadsheetImage
from PIL import Image
from pptx import Presentation
from pptx.util import Inches
import pytest

from file2doc.app import create_app
from file2doc.parsers import ParseOptions
from fixtures import sample_file


class _Responses:
    def create(self, **kwargs):
        return SimpleNamespace(
            output_text=json.dumps(
                {
                    "description": "An embedded document image.",
                    "visibleText": [],
                    "candidateNumericValues": [],
                    "layout": "Embedded in document reading order.",
                    "warnings": [],
                }
            )
        )


class _FailingResponses:
    def create(self, **kwargs):
        raise RuntimeError("provider unavailable")


def _visual_parse_options() -> ParseOptions:
    return ParseOptions(
        visual_client=SimpleNamespace(responses=_Responses()),
        visual_model="ep-visual",
    )


@pytest.mark.parametrize(
    ("filename", "content_type", "source_bytes", "native_marker"),
    [
        (
            "partial.docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            lambda: _docx_with_failed_visual(),
            "DOCX native content survives",
        ),
        (
            "partial.pptx",
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            lambda: _pptx_with_failed_visual(),
            "PPTX native content survives",
        ),
        (
            "partial.xlsx",
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            lambda: _xlsx_with_failed_visual(),
            "XLSX native content survives",
        ),
    ],
)
def test_uploaded_office_with_failed_visual_completes_with_warnings(
    tmp_path,
    filename,
    content_type,
    source_bytes,
    native_marker,
):
    client = TestClient(
        create_app(
            storage_root=tmp_path / "storage",
            auth_enabled=False,
            parse_options=ParseOptions(
                visual_client=SimpleNamespace(responses=_FailingResponses()),
                visual_model="ep-visual",
            ),
        )
    )

    created = client.post(
        "/parse-jobs/upload",
        files={"file": (filename, source_bytes(), content_type)},
    ).json()
    job = client.get(created["poll_url"]).json()

    assert job["status"] == "completed_with_warnings"
    assert job["warnings_count"] == 1
    manifest = client.get(job["result"]["manifest_url"]).json()
    assert manifest["warnings"][0]["code"] == "image_parse_warning"
    assert "provider unavailable" in manifest["warnings"][0]["message"]
    markdown = client.get(job["result"]["content_url"]).text
    assert native_marker in markdown
    assert "## Warnings" in markdown
    assert "provider unavailable" in markdown


def test_uploaded_docx_produces_non_empty_markitdown_markdown(tmp_path):
    sample = tmp_path / "office-support.docx"
    document = Document()
    document.add_heading("Quarterly Office Support", level=1)
    document.add_paragraph("DOCX generated sample marker: alpha roadmap.")
    document.save(sample)
    client = TestClient(
        create_app(
            storage_root=tmp_path / "storage",
            auth_enabled=False,
            parse_options=_visual_parse_options(),
        )
    )

    create_response = client.post(
        "/parse-jobs/upload",
        files={
            "file": (
                sample.name,
                sample.read_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"parser_profile": "agent", "retention": "short"},
    )

    assert create_response.status_code == 201
    created = create_response.json()

    poll_response = client.get(created["poll_url"])
    assert poll_response.status_code == 200
    job = poll_response.json()
    assert job["status"] == "completed"
    assert job["stage"] == "completed"
    assert job["percent"] == 100

    manifest_response = client.get(job["result"]["manifest_url"])
    assert manifest_response.status_code == 200
    manifest = manifest_response.json()
    assert manifest["source"]["filename"] == sample.name
    assert manifest["source"]["content_type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert manifest["content"]["artifact_id"] == job["result"]["content_artifact_id"]
    assert manifest["parser"]["name"] == "file2doc-markitdown-visual"
    assert manifest["parser"]["version"]
    assert manifest["parser"]["visual_plugin_version"] == "0.3.2"

    content_response = client.get(job["result"]["content_url"])
    assert content_response.status_code == 200
    assert content_response.headers["content-type"].startswith("text/markdown")
    content = content_response.text
    assert "Quarterly Office Support" in content
    assert "DOCX generated sample marker: alpha roadmap." in content


def test_uploaded_xlsx_produces_non_empty_markitdown_markdown(tmp_path):
    sample = tmp_path / "office-support.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Summary"
    worksheet.append(["Metric", "Value"])
    worksheet.append(["XLSX generated sample marker", "beta spreadsheet"])
    workbook.save(sample)
    client = TestClient(
        create_app(
            storage_root=tmp_path / "storage",
            auth_enabled=False,
            parse_options=_visual_parse_options(),
        )
    )

    create_response = client.post(
        "/parse-jobs/upload",
        files={
            "file": (
                sample.name,
                sample.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
        data={"parser_profile": "agent", "retention": "short"},
    )

    assert create_response.status_code == 201
    created = create_response.json()

    poll_response = client.get(created["poll_url"])
    assert poll_response.status_code == 200
    job = poll_response.json()
    assert job["status"] == "completed"
    assert job["stage"] == "completed"
    assert job["percent"] == 100

    manifest_response = client.get(job["result"]["manifest_url"])
    assert manifest_response.status_code == 200
    manifest = manifest_response.json()
    assert manifest["source"]["filename"] == sample.name
    assert manifest["source"]["content_type"] == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert manifest["content"]["artifact_id"] == job["result"]["content_artifact_id"]
    assert manifest["parser"]["name"] == "file2doc-markitdown-visual"
    assert manifest["parser"]["version"]
    assert manifest["parser"]["visual_plugin_version"] == "0.3.2"

    content_response = client.get(job["result"]["content_url"])
    assert content_response.status_code == 200
    assert content_response.headers["content-type"].startswith("text/markdown")
    content = content_response.text
    assert "Summary" in content
    assert "XLSX generated sample marker" in content
    assert "beta spreadsheet" in content


def test_uploaded_docx_with_no_extracted_markdown_fails_visibly(tmp_path):
    sample = tmp_path / "blank-office-support.docx"
    Document().save(sample)
    client = TestClient(
        create_app(
            storage_root=tmp_path / "storage",
            auth_enabled=False,
            parse_options=_visual_parse_options(),
        )
    )

    created = client.post(
        "/parse-jobs/upload",
        files={
            "file": (
                sample.name,
                sample.read_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    ).json()
    job = client.get(created["poll_url"]).json()

    assert job["status"] == "failed"
    assert job["stage"] == "failed"
    assert job["error"]["code"] == "visual_item_failed"
    assert job["result"] is None

    manifest_response = client.get(f"/parse-jobs/{created['job_id']}/result")
    assert manifest_response.status_code == 409
    assert manifest_response.json()["detail"]["code"] == "result_not_ready"


def test_uploaded_pptx_produces_non_empty_markitdown_markdown(tmp_path):
    sample = sample_file(
        "1.1.1 基础系列-导读课-大模型技术趋势与企业级 LLMOps 平台价值解读.pptx"
    )
    client = TestClient(
        create_app(
            storage_root=tmp_path,
            auth_enabled=False,
            parse_options=_visual_parse_options(),
        )
    )

    create_response = client.post(
        "/parse-jobs/upload",
        files={
            "file": (
                sample.name,
                sample.read_bytes(),
                "application/vnd.openxmlformats-officedocument.presentationml.presentation",
            )
        },
        data={"parser_profile": "agent", "retention": "short"},
    )

    assert create_response.status_code == 201
    created = create_response.json()

    poll_response = client.get(created["poll_url"])
    assert poll_response.status_code == 200
    job = poll_response.json()
    assert job["status"] == "completed"
    assert job["stage"] == "completed"
    assert job["percent"] == 100

    manifest_response = client.get(job["result"]["manifest_url"])
    assert manifest_response.status_code == 200
    manifest = manifest_response.json()
    assert manifest["source"]["filename"] == sample.name
    assert manifest["source"]["content_type"] == (
        "application/vnd.openxmlformats-officedocument.presentationml.presentation"
    )
    assert manifest["content"]["artifact_id"] == job["result"]["content_artifact_id"]

    content_response = client.get(job["result"]["content_url"])
    assert content_response.status_code == 200
    assert content_response.headers["content-type"].startswith("text/markdown")
    content = content_response.text
    assert "大模型技术趋势与企业级 LLMOps 平台价值解读" in content
    assert "AI - 你的必备「数字生产力」" in content
    assert len(content) > 1000


def _png_bytes() -> bytes:
    image = Image.new("RGB", (24, 16), color=(32, 64, 192))
    output = io.BytesIO()
    image.save(output, format="PNG")
    return output.getvalue()


def _docx_with_failed_visual() -> bytes:
    document = Document()
    document.add_paragraph("DOCX native content survives")
    document.add_picture(io.BytesIO(_png_bytes()))
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pptx_with_failed_visual() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(
        Inches(1), Inches(1), Inches(5), Inches(0.5)
    ).text = "PPTX native content survives"
    slide.shapes.add_picture(io.BytesIO(_png_bytes()), Inches(1), Inches(2))
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def _xlsx_with_failed_visual() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["XLSX native content survives"])
    sheet.add_image(SpreadsheetImage(io.BytesIO(_png_bytes())), "A3")
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
