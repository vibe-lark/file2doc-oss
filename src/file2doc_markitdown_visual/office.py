"""Office converters with inline visual parsing for embedded images.

Derived from Microsoft ``markitdown-ocr`` at commit
``e144e0a2be95b34df17433bac904e635f2c5e551``, especially its DOCX, PPTX,
and XLSX converter modules. The MIT attribution is retained in NOTICE.md.
"""

from __future__ import annotations

import io
import re
from operator import attrgetter
from typing import Any, BinaryIO
from uuid import uuid4

import mammoth
import pandas as pd
import pptx
from markitdown import DocumentConverterResult, StreamInfo
from markitdown.converter_utils.docx.pre_process import pre_process_docx
from markitdown.converters import HtmlConverter, PptxConverter
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .embedded import EmbeddedVisualParser, render_embedded_visual


_DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)
_PPTX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)
_XLSX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class VisualDocxConverter(HtmlConverter):
    def __init__(self, *, visual_parser: EmbeddedVisualParser) -> None:
        super().__init__()
        self._html_converter = HtmlConverter()
        self._visual_parser = visual_parser

    def accepts(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,
    ) -> bool:
        mimetype = (stream_info.mimetype or "").lower()
        extension = (stream_info.extension or "").lower()
        return extension == ".docx" or mimetype == _DOCX_MIME_TYPE

    def convert(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,
    ) -> DocumentConverterResult:
        visual_session = self._visual_parser.new_session()
        visual_blocks: list[str] = []
        visual_warnings: list[str] = []
        placeholder_prefix = f"FILE2DOCVISUAL-{uuid4().hex}-"
        placeholders: list[str] = []

        def convert_image(image):
            location = f"DOCX image {len(visual_blocks) + 1}"
            with image.open() as image_stream:
                result = visual_session.parse(
                    image_stream.read(),
                    content_type=image.content_type,
                    location=location,
                )
            visual_blocks.append(render_embedded_visual(result, location=location))
            if result.warning is not None:
                visual_warnings.append(f"{location}: {result.warning}")
            placeholder = f"{placeholder_prefix}{len(visual_blocks) - 1}"
            placeholders.append(placeholder)
            return {"src": placeholder}

        file_stream.seek(0)
        processed = pre_process_docx(file_stream)
        html = mammoth.convert_to_html(
            processed,
            style_map=kwargs.get("style_map"),
            convert_image=mammoth.images.img_element(convert_image),
        ).value
        html = _replace_docx_images_with_placeholders(html, placeholders)
        markdown = self._html_converter.convert_string(html, **kwargs).markdown
        for placeholder, block in zip(placeholders, visual_blocks):
            markdown = markdown.replace(placeholder, block, 1)
        markdown = _append_visual_warnings(markdown, visual_warnings)
        return DocumentConverterResult(markdown=markdown)


def _replace_docx_images_with_placeholders(
    html: str,
    placeholders: list[str],
) -> str:
    for placeholder in placeholders:
        html = re.sub(
            rf'<img\b[^>]*\bsrc=["\']{re.escape(placeholder)}["\'][^>]*>',
            f"<p>{placeholder}</p>",
            html,
            count=1,
        )
    return html


class VisualPptxConverter(PptxConverter):
    """Preserve MarkItDown's slide/shape conversion and parse every picture."""

    def __init__(self, *, visual_parser: EmbeddedVisualParser) -> None:
        super().__init__()
        self._visual_parser = visual_parser

    def accepts(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,
    ) -> bool:
        mimetype = (stream_info.mimetype or "").lower()
        extension = (stream_info.extension or "").lower()
        return extension == ".pptx" or mimetype == _PPTX_MIME_TYPE

    def convert(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,
    ) -> DocumentConverterResult:
        visual_session = self._visual_parser.new_session()
        presentation = pptx.Presentation(file_stream)
        output: list[str] = []
        visual_warnings: list[str] = []

        for slide_number, slide in enumerate(presentation.slides, start=1):
            output.extend(["", f"<!-- Slide number: {slide_number} -->"])
            title = slide.shapes.title

            def append_shape(shape) -> None:
                if self._is_picture(shape):
                    result = visual_session.parse(
                        shape.image.blob,
                        content_type=shape.image.content_type,
                        location=f"PPTX slide {slide_number}, shape {shape.name}",
                    )
                    location = f"PPTX slide {slide_number}, shape {shape.name}"
                    output.append(render_embedded_visual(result, location=location))
                    if result.warning is not None:
                        visual_warnings.append(f"{location}: {result.warning}")

                if self._is_table(shape):
                    output.append(
                        self._convert_table_to_markdown(shape.table, **kwargs).strip()
                    )

                if shape.has_chart:
                    output.append(self._convert_chart_to_markdown(shape.chart).strip())
                elif shape.has_text_frame:
                    text = shape.text.lstrip() if shape == title else shape.text
                    output.append(f"# {text}" if shape == title else text)

                if shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
                    for child in sorted(shape.shapes, key=attrgetter("top", "left")):
                        append_shape(child)

            for shape in sorted(slide.shapes, key=attrgetter("top", "left")):
                append_shape(shape)

            if slide.has_notes_slide:
                notes_frame = slide.notes_slide.notes_text_frame
                if notes_frame is not None and notes_frame.text.strip():
                    output.extend(["### Notes:", notes_frame.text])

        markdown = "\n\n".join(part.strip() for part in output if part.strip())
        return DocumentConverterResult(
            markdown=_append_visual_warnings(markdown, visual_warnings)
        )


class VisualXlsxConverter(HtmlConverter):
    """Preserve sheet tables and append anchored visual evidence per sheet."""

    def __init__(self, *, visual_parser: EmbeddedVisualParser) -> None:
        super().__init__()
        self._html_converter = HtmlConverter()
        self._visual_parser = visual_parser

    def accepts(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,
    ) -> bool:
        mimetype = (stream_info.mimetype or "").lower()
        extension = (stream_info.extension or "").lower()
        return extension == ".xlsx" or mimetype == _XLSX_MIME_TYPE

    def convert(
        self,
        file_stream: BinaryIO,
        stream_info: StreamInfo,
        **kwargs: Any,
    ) -> DocumentConverterResult:
        visual_session = self._visual_parser.new_session()
        source = file_stream.read()
        workbook = load_workbook(io.BytesIO(source))
        sheets = pd.read_excel(
            io.BytesIO(source),
            sheet_name=None,
            engine="openpyxl",
        )
        output: list[str] = []
        visual_warnings: list[str] = []

        for sheet_name in workbook.sheetnames:
            output.append(f"## {sheet_name}")
            frame = sheets[sheet_name]
            html = frame.to_html(index=False)
            output.append(
                self._html_converter.convert_string(html, **kwargs).markdown.strip()
            )

            sheet = workbook[sheet_name]
            images = sorted(
                enumerate(getattr(sheet, "_images", [])),
                key=lambda item: (*_xlsx_anchor_coordinates(item[1]), item[0]),
            )
            if images:
                output.append("### Embedded Images")
            for _, image in images:
                row, column = _xlsx_anchor_coordinates(image)
                cell = f"{get_column_letter(column + 1)}{row + 1}"
                result = visual_session.parse(
                    image._data(),
                    content_type=_xlsx_content_type(image),
                    location=f"XLSX sheet {sheet_name}, cell {cell}",
                )
                location = f"XLSX sheet {sheet_name}, cell {cell}"
                output.append(render_embedded_visual(result, location=location))
                if result.warning is not None:
                    visual_warnings.append(f"{location}: {result.warning}")

        markdown = "\n\n".join(output).strip()
        return DocumentConverterResult(
            markdown=_append_visual_warnings(markdown, visual_warnings)
        )


def _append_visual_warnings(markdown: str, warnings: list[str]) -> str:
    if not warnings:
        return markdown
    warning_section = "\n".join(
        ["## Warnings", "", *(f"- {item}" for item in warnings)]
    )
    return f"{markdown.rstrip()}\n\n{warning_section}\n"


def _xlsx_anchor_coordinates(image) -> tuple[int, int]:
    anchor = getattr(image, "anchor", None)
    marker = getattr(anchor, "_from", None)
    return (getattr(marker, "row", 0), getattr(marker, "col", 0))


def _xlsx_content_type(image) -> str:
    image_format = (getattr(image, "format", None) or "png").lower()
    if image_format in {"jpg", "jpeg"}:
        return "image/jpeg"
    return f"image/{image_format}"
