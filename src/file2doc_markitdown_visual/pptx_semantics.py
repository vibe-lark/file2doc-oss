from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


@dataclass(frozen=True)
class ExcludedPptxShape:
    source_slide_identity: str
    native_slide_index: int
    shape_name: str


@dataclass(frozen=True)
class PptxDynamicField:
    field_type: str
    cached_value: str
    shape_name: str
    left: int
    top: int
    width: int
    height: int


@dataclass(frozen=True)
class PptxCoordinateTransform:
    scale_x: float = 1.0
    scale_y: float = 1.0
    offset_x: float = 0.0
    offset_y: float = 0.0


def shape_intersects_slide(
    shape,
    *,
    slide_width: int,
    slide_height: int,
    transform: PptxCoordinateTransform = PptxCoordinateTransform(),
) -> bool:
    left, top, width, height = shape_bounds_in_slide(shape, transform=transform)
    right = left + width
    bottom = top + height
    return right > 0 and bottom > 0 and left < slide_width and top < slide_height


def shape_bounds_in_slide(
    shape,
    *,
    transform: PptxCoordinateTransform,
) -> tuple[float, float, float, float]:
    return (
        transform.offset_x + int(shape.left) * transform.scale_x,
        transform.offset_y + int(shape.top) * transform.scale_y,
        int(shape.width) * transform.scale_x,
        int(shape.height) * transform.scale_y,
    )


def group_child_transform(
    group,
    *,
    parent: PptxCoordinateTransform,
) -> PptxCoordinateTransform:
    group_left, group_top, _group_width, _group_height = shape_bounds_in_slide(
        group,
        transform=parent,
    )
    xfrm = group._element.grpSpPr.xfrm
    child_extent_x = int(xfrm.chExt.cx)
    child_extent_y = int(xfrm.chExt.cy)
    scale_x = parent.scale_x * (
        int(group.width) / child_extent_x if child_extent_x else 1.0
    )
    scale_y = parent.scale_y * (
        int(group.height) / child_extent_y if child_extent_y else 1.0
    )
    return PptxCoordinateTransform(
        scale_x=scale_x,
        scale_y=scale_y,
        offset_x=group_left - int(xfrm.chOff.x) * scale_x,
        offset_y=group_top - int(xfrm.chOff.y) * scale_y,
    )


def inspect_off_canvas_shapes(source_path: Path) -> tuple[ExcludedPptxShape, ...]:
    presentation = Presentation(source_path)
    excluded: list[ExcludedPptxShape] = []
    for native_slide_index, slide in enumerate(presentation.slides, 1):
        def inspect(shape, transform: PptxCoordinateTransform) -> None:
            if shape_intersects_slide(
                shape,
                slide_width=presentation.slide_width,
                slide_height=presentation.slide_height,
                transform=transform,
            ):
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    child_transform = group_child_transform(shape, parent=transform)
                    for child in shape.shapes:
                        inspect(child, child_transform)
                return
            excluded.append(
                ExcludedPptxShape(
                    source_slide_identity=f"pptx-slide-{slide.slide_id}",
                    native_slide_index=native_slide_index,
                    shape_name=shape.name,
                )
            )

        for shape in slide.shapes:
            inspect(shape, PptxCoordinateTransform())
    return tuple(excluded)


def dynamic_fields_for_slide(slide) -> tuple[PptxDynamicField, ...]:
    return tuple(
        field
        for shape in slide.shapes
        for field in dynamic_fields_for_shape(shape)
    )


def shape_text_without_dynamic_fields(shape) -> str:
    paragraphs: list[str] = []
    for paragraph in shape.text_frame.paragraphs:
        parts: list[str] = []
        for child in paragraph._p:
            if child.tag == qn("a:r"):
                text = child.find(qn("a:t"))
                if text is not None and text.text:
                    parts.append(text.text)
            elif child.tag == qn("a:br"):
                parts.append("\v")
        paragraphs.append("".join(parts))
    return "\n".join(paragraphs)


def dynamic_fields_for_shape(shape) -> tuple[PptxDynamicField, ...]:
    fields: list[PptxDynamicField] = []
    for field in shape._element.iter(qn("a:fld")):
        text = field.find(qn("a:t"))
        fields.append(
            PptxDynamicField(
                field_type=str(field.get("type") or "unknown"),
                cached_value=(text.text if text is not None and text.text else ""),
                shape_name=shape.name,
                left=int(shape.left),
                top=int(shape.top),
                width=int(shape.width),
                height=int(shape.height),
            )
        )
    return tuple(fields)
