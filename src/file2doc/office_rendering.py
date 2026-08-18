from __future__ import annotations

import gc
import subprocess
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable
from uuid import uuid4

import pypdfium2 as pdfium
from PIL import Image
from pptx import Presentation
from pptx.util import Pt

from file2doc.rendering import AGENT_PAGE_IMAGE_DPI, AGENT_THUMBNAIL_MAX_EDGE
from file2doc_markitdown_visual.pptx_semantics import (
    PptxDynamicField,
    dynamic_fields_for_slide,
)


PPTX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.presentationml.presentation"
)


class OfficeRenderError(RuntimeError):
    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class PptxSlide:
    source_slide_identity: str
    native_slide_index: int
    hidden: bool
    dynamic_fields: tuple[PptxDynamicField, ...]

    def source_ref(self, *, rendered_page_index: int | None) -> dict:
        return {
            "type": "pptx_slide",
            "source_slide_identity": self.source_slide_identity,
            "native_slide_index": self.native_slide_index,
            "hidden": self.hidden,
            "rendered_page_index": rendered_page_index,
        }


def is_pptx_source(source_path: Path, content_type: str) -> bool:
    return (
        content_type.split(";", 1)[0] == PPTX_CONTENT_TYPE
        or source_path.suffix.lower() == ".pptx"
    )


def inspect_pptx_slides(source_path: Path) -> list[PptxSlide]:
    slides, _width, _height = _inspect_pptx_source(source_path)
    return slides


def _inspect_pptx_source(source_path: Path) -> tuple[list[PptxSlide], int, int]:
    presentation = Presentation(source_path)
    slides = _slides_from_presentation(presentation)
    width = presentation.slide_width
    height = presentation.slide_height
    del presentation
    gc.collect()
    return slides, width, height


def _slides_from_presentation(presentation) -> list[PptxSlide]:
    return [
        PptxSlide(
            source_slide_identity=f"pptx-slide-{slide.slide_id}",
            native_slide_index=index,
            hidden=_is_hidden(slide),
            dynamic_fields=dynamic_fields_for_slide(slide),
        )
        for index, slide in enumerate(presentation.slides, 1)
    ]


def render_pptx_visual_assets(
    source_path: Path,
    result_root: Path,
    *,
    new_artifact_id: Callable[[], str],
    page_image_dpi: int = AGENT_PAGE_IMAGE_DPI,
    thumbnail_max_edge: int = AGENT_THUMBNAIL_MAX_EDGE,
) -> tuple[list[dict], list[dict], dict[str, dict]]:
    slides, presentation_width, presentation_height = _inspect_pptx_source(
        source_path
    )
    visible_slides = [slide for slide in slides if not slide.hidden]
    with tempfile.TemporaryDirectory(prefix="file2doc-pptx-render-") as tmpdir:
        temporary_root = Path(tmpdir)
        marked_root = temporary_root / "binding"
        marked_source, marker_by_identity = _instrumented_copy(
            source_path,
            marked_root,
            slides,
        )
        marked_pdf = _convert_to_pdf(marked_source, marked_root)
        rendered_page_by_identity = _verified_render_binding(
            marked_pdf,
            marker_by_identity,
            visible_slides,
        )
        clean_root = temporary_root / "clean"
        pdf_path = _convert_to_pdf(source_path, clean_root)
        document = pdfium.PdfDocument(pdf_path)
        rendered_page_count = len(document)
        if rendered_page_count != len(rendered_page_by_identity):
            document.close()
            raise OfficeRenderError(
                "pptx_slide_binding_unprovable",
                "Cannot prove PPTX slide binding: "
                f"{len(rendered_page_by_identity)} verified native slides produced "
                f"{rendered_page_count} rendered pages",
            )
        page_index: list[dict] = []
        media_index: list[dict] = []
        artifacts: dict[str, dict] = {}
        for slide in slides:
            rendered_page_index = rendered_page_by_identity.get(
                slide.source_slide_identity
            )
            entry = {
                "source_slide_identity": slide.source_slide_identity,
                "source_unit": "slide",
                "source_unit_index": slide.native_slide_index,
                "native_slide_index": slide.native_slide_index,
                "hidden": slide.hidden,
                "rendered_page_index": rendered_page_index,
                "render_binding": {
                    "status": (
                        "verified"
                        if rendered_page_index is not None
                        else "not_rendered"
                    ),
                    "method": "instrumented_render_marker",
                },
                "parse_status": "parsed",
            }
            if slide.dynamic_fields:
                entry["dynamic_fields"] = _dynamic_field_documents(
                    slide,
                    document[rendered_page_index - 1]
                    if rendered_page_index is not None
                    else None,
                    presentation_width=presentation_width,
                    presentation_height=presentation_height,
                )
            if rendered_page_index is None:
                page_index.append(entry)
                continue

            image = document[rendered_page_index - 1].render(
                scale=page_image_dpi / 72
            ).to_pil()
            page_image_path = Path(
                f"images/pages/slide_{slide.native_slide_index:03d}.png"
            )
            thumbnail_path = Path(
                f"images/thumbs/slide_{slide.native_slide_index:03d}.jpg"
            )
            _save_image(image, result_root / page_image_path)
            thumbnail = image.copy()
            thumbnail.thumbnail((thumbnail_max_edge, thumbnail_max_edge))
            if thumbnail.mode != "RGB":
                thumbnail = thumbnail.convert("RGB")
            _save_image(thumbnail, result_root / thumbnail_path, format="JPEG")

            image_artifact_id = new_artifact_id()
            thumbnail_artifact_id = new_artifact_id()
            image_id = f"{slide.source_slide_identity}-image"
            thumbnail_id = f"{slide.source_slide_identity}-thumb"
            source_ref = slide.source_ref(
                rendered_page_index=rendered_page_index
            )
            media_index.extend(
                [
                    {
                        "id": image_id,
                        "kind": "page_image",
                        "page": rendered_page_index,
                        "path": page_image_path.as_posix(),
                        "artifact_id": image_artifact_id,
                        "media_type": "image/png",
                        "thumbnail_id": thumbnail_id,
                        "source_ref": source_ref,
                        "derived": False,
                        "derivation": {"renderer": "LibreOffice+PDFium"},
                    },
                    {
                        "id": thumbnail_id,
                        "kind": "thumbnail",
                        "page": rendered_page_index,
                        "path": thumbnail_path.as_posix(),
                        "artifact_id": thumbnail_artifact_id,
                        "media_type": "image/jpeg",
                        "source_ref": source_ref,
                        "derived": False,
                        "derivation": {"renderer": "LibreOffice+PDFium"},
                    },
                ]
            )
            artifacts.update(
                {
                    image_artifact_id: {
                        "artifact_id": image_artifact_id,
                        "kind": "page_image",
                        "path": page_image_path.as_posix(),
                        "media_type": "image/png",
                        "source_ref": source_ref,
                    },
                    thumbnail_artifact_id: {
                        "artifact_id": thumbnail_artifact_id,
                        "kind": "thumbnail",
                        "path": thumbnail_path.as_posix(),
                        "media_type": "image/jpeg",
                        "source_ref": source_ref,
                    },
                }
            )
            page_index.append(
                entry
                | {
                    "page_image_id": image_id,
                    "thumbnail_id": thumbnail_id,
                }
            )
            thumbnail.close()
            image.close()
        document.close()
        return page_index, media_index, artifacts


def _convert_to_pdf(source_path: Path, temporary_root: Path) -> Path:
    temporary_root.mkdir(parents=True, exist_ok=True)
    profile = temporary_root / "libreoffice-profile"
    command = [
        "libreoffice",
        "--headless",
        f"-env:UserInstallation={profile.as_uri()}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(temporary_root),
        str(source_path),
    ]
    try:
        completed = subprocess.run(
            command,
            check=False,
            capture_output=True,
            text=True,
            timeout=300,
        )
    except (FileNotFoundError, subprocess.TimeoutExpired) as error:
        raise OfficeRenderError(
            "pptx_renderer_unavailable",
            f"LibreOffice could not render the PPTX source: {error}",
        ) from error
    pdf_path = temporary_root / f"{source_path.stem}.pdf"
    if completed.returncode != 0 or not pdf_path.is_file():
        detail = (completed.stderr or completed.stdout).strip()
        raise OfficeRenderError(
            "pptx_render_failed",
            "LibreOffice did not produce a complete PPTX render"
            + (f": {detail}" if detail else ""),
        )
    return pdf_path


def _instrumented_copy(
    source_path: Path,
    output_root: Path,
    slides: list[PptxSlide],
) -> tuple[Path, dict[str, str]]:
    output_root.mkdir(parents=True, exist_ok=True)
    presentation = Presentation(source_path)
    nonce = uuid4().hex[:12]
    marker_by_identity = {
        slide.source_slide_identity: f"F2D-{nonce}-{slide.source_slide_identity}"
        for slide in slides
    }
    for slide_document, slide in zip(presentation.slides, slides):
        marker = marker_by_identity[slide.source_slide_identity]
        text_box = slide_document.shapes.add_textbox(
            0,
            0,
            presentation.slide_width,
            250_000,
        )
        text_box.name = f"File2Doc binding marker {slide.source_slide_identity}"
        text_box.text = marker
        for paragraph in text_box.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(6)
    marked_source = output_root / "file2doc-binding.pptx"
    presentation.save(marked_source)
    del presentation
    gc.collect()
    return marked_source, marker_by_identity


def _verified_render_binding(
    marked_pdf: Path,
    marker_by_identity: dict[str, str],
    visible_slides: list[PptxSlide],
) -> dict[str, int]:
    identity_by_marker = {
        marker: identity for identity, marker in marker_by_identity.items()
    }
    expected = {slide.source_slide_identity for slide in visible_slides}
    document = pdfium.PdfDocument(marked_pdf)
    found: dict[str, int] = {}
    try:
        for rendered_page_index in range(1, len(document) + 1):
            text_page = document[rendered_page_index - 1].get_textpage()
            try:
                page_text = text_page.get_text_range()
            finally:
                text_page.close()
            normalized_page_text = "".join(page_text.split())
            matches = [
                identity
                for marker, identity in identity_by_marker.items()
                if marker in normalized_page_text
            ]
            if len(matches) != 1 or matches[0] in found:
                raise OfficeRenderError(
                    "pptx_slide_binding_unprovable",
                    "Instrumented PPTX render did not expose exactly one unique "
                    f"Source slide identity on rendered page {rendered_page_index}",
                )
            found[matches[0]] = rendered_page_index
    finally:
        document.close()
    if set(found) != expected:
        raise OfficeRenderError(
            "pptx_slide_binding_unprovable",
            "Instrumented PPTX render identities do not match visible native slides",
        )
    return found


def _is_hidden(slide) -> bool:
    return str(slide._element.get("show", "1")).lower() in {"0", "false"}


def _dynamic_field_documents(
    slide: PptxSlide,
    rendered_page,
    *,
    presentation_width: int,
    presentation_height: int,
) -> list[dict]:
    documents = []
    for field in slide.dynamic_fields:
        evaluated_value = _rendered_field_value(
            field,
            rendered_page,
            presentation_width=presentation_width,
            presentation_height=presentation_height,
        )
        documents.append(
            {
                "field_type": field.field_type,
                "shape_name": field.shape_name,
                "cached_value": field.cached_value,
                "cached_value_authority": "non_authoritative",
                "evaluated_value": evaluated_value,
                "evaluated_value_authority": (
                    "canonical_render" if evaluated_value else "unavailable"
                ),
                "authoritative_representation": (
                    "evaluated_value" if evaluated_value else "page_image"
                ),
            }
        )
    return documents


def _rendered_field_value(
    field: PptxDynamicField,
    rendered_page,
    *,
    presentation_width: int,
    presentation_height: int,
) -> str | None:
    if rendered_page is None:
        return None
    page_width, page_height = rendered_page.get_size()
    left = field.left / presentation_width * page_width
    right = (field.left + field.width) / presentation_width * page_width
    bottom = page_height - (
        (field.top + field.height) / presentation_height * page_height
    )
    top = page_height - (field.top / presentation_height * page_height)
    text_page = rendered_page.get_textpage()
    try:
        value = text_page.get_text_bounded(left, bottom, right, top).strip()
    finally:
        text_page.close()
    return value or None


def _save_image(image: Image.Image, path: Path, *, format: str | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    image.save(path, format=format)
